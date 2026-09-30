import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    PRIMARY = RGBColor(15, 138, 95)      # #0F8A5F Emerald Green
    PRIMARY_DARK = RGBColor(11, 107, 73) # #0B6B49
    SECONDARY = RGBColor(70, 179, 106)   # #46B36A
    ACCENT = RGBColor(246, 183, 60)      # #F6B73C Warm Gold
    DARK_BG = RGBColor(15, 23, 42)       # #0F172A Dark Navy
    CARD_BG = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(248, 250, 252)   # #F8FAFC
    TEXT_MAIN = RGBColor(30, 41, 59)     # #1E293B
    TEXT_MUTED = RGBColor(100, 116, 139) # #64748B
    BORDER_COLOR = RGBColor(226, 232, 240)

    def set_slide_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    def add_header(slide, tag, title, subtitle=None, dark=False):
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = ACCENT if dark else PRIMARY

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(255, 255, 255) if dark else TEXT_MAIN

        if subtitle:
            p_sub = tf_title.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = RGBColor(148, 163, 184) if dark else TEXT_MUTED

    # =========================================================================
    # SLIDE 1: TITLE & HERO
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide1, DARK_BG)

    # Accent glow top border
    glow = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
    glow.fill.solid()
    glow.fill.fore_color.rgb = PRIMARY
    glow.line.fill.background()

    # Tag Badge
    badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(3.4), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(20, 83, 45)
    badge.line.fill.background()
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "🌱 NEXT-GEN DAIRY FINTECH"
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = RGBColor(134, 239, 172)
    p_b.alignment = PP_ALIGN.CENTER

    # Main Heading
    h_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.5), Inches(2.2))
    tf_h = h_box.text_frame
    tf_h.word_wrap = True
    p_h1 = tf_h.paragraphs[0]
    p_h1.text = "3T (Time To Time) Dairy"
    p_h1.font.size = Pt(40)
    p_h1.font.bold = True
    p_h1.font.color.rgb = RGBColor(255, 255, 255)

    p_h2 = tf_h.add_paragraph()
    p_h2.text = "Milk Now. Money Now."
    p_h2.font.size = Pt(28)
    p_h2.font.bold = True
    p_h2.font.color.rgb = ACCENT

    # Description
    desc_box = slide1.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(8.0), Inches(1.2))
    tf_d = desc_box.text_frame
    tf_d.word_wrap = True
    p_d = tf_d.paragraphs[0]
    p_d.text = "India's first IoT-connected dairy terminal delivering automated quality analysis, transparent dual displays, and instant 2-second direct-to-bank UPI payouts for 80M+ farmers."
    p_d.font.size = Pt(14)
    p_d.font.color.rgb = RGBColor(203, 213, 225)

    # 3 Stat Cards on Right
    stats = [
        ("₹10,000 Cr+", "Annual Delayed Farmer Working Capital"),
        ("2 Seconds", "Direct UPI / Bank Payout per Milk Delivery"),
        ("100% Zero-Touch", "Bluetooth Scale + Dual Display Anti-Tamper")
    ]
    for idx, (val, lbl) in enumerate(stats):
        top_pos = 1.6 + idx * 1.7
        card = add_card(slide1, 9.2, top_pos, 3.3, 1.4, bg_color=RGBColor(30, 41, 59), border_color=RGBColor(51, 65, 85))
        tf_c = card.text_frame
        tf_c.word_wrap = True
        p_v = tf_c.paragraphs[0]
        p_v.text = val
        p_v.font.size = Pt(22)
        p_v.font.bold = True
        p_v.font.color.rgb = RGBColor(134, 239, 172)
        p_l = tf_c.add_paragraph()
        p_l.text = lbl
        p_l.font.size = Pt(10)
        p_l.font.color.rgb = RGBColor(148, 163, 184)

    # Footer note
    f_box = slide1.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Founder: Sandesh Kadam | 3T Dairy Fintech Innovation Platform"
    p_f.font.size = Pt(11)
    p_f.font.color.rgb = RGBColor(100, 116, 139)

    # =========================================================================
    # SLIDE 2: THE PROBLEM
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide2, LIGHT_BG)
    add_header(slide2, "THE GROUND REALITY", "3 Critical Pain Points of India's 80 Million Dairy Farmers", "Why traditional paper registers & delayed cooperative credit cycles are hurting rural livelihood")

    problems = [
        ("⏳ 15–20 Day Payment Delay", 
         "Farmers incur daily expenses for cattle feed and medicine, but dairy payments are held for weeks. This forces farmers into high-interest informal loans.",
         "Impact: Severe rural working-capital crisis"),
        ("❌ Opacity & Manual Tampering", 
         "Milk weight and FAT/SNF quality are written manually into registers. Operators frequently under-read quality metrics, reducing farmer earnings by 10–18%.",
         "Impact: Lack of trust between farmers & centers"),
        ("📉 Zero Early Health Warning", 
         "When cattle fall ill, milk yield drops silently. Farmers only realize days later when infection spreads, causing permanent loss in cattle productivity.",
         "Impact: ₹4,000+ loss per diseased cattle")
    ]

    for idx, (title, desc, impact) in enumerate(problems):
        left_pos = 0.8 + idx * 3.95
        card = add_card(slide2, left_pos, 1.9, 3.8, 4.8)
        tf = card.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(185, 28, 28)

        p2 = tf.add_paragraph()
        p2.text = "\n" + desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MAIN

        p3 = tf.add_paragraph()
        p3.text = "\n\n🚨 " + impact
        p3.font.size = Pt(11)
        p3.font.bold = True
        p3.font.color.rgb = RGBColor(220, 38, 38)

    # =========================================================================
    # SLIDE 3: THE SOLUTION (4-PILLAR ECOSYSTEM)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide3, LIGHT_BG)
    add_header(slide3, "THE 3T BREAKTHROUGH", "A Unified IoT & Instant Payout Ecosystem", "Transforming village milk collection into a transparent, high-speed digital experience")

    pillars = [
        ("⚖️ Web Bluetooth IoT Scale", "Captures milk volume directly from digital scales via Web Bluetooth API (BLE 4.0+) with zero human typing.", PRIMARY),
        ("🖥️ Customer Dual Display", "Live wireless synchronization via BroadcastChannel API so farmers standing in queue verify their FAT/SNF in real time.", PRIMARY_DARK),
        ("⚡ Instant 2-Sec UPI Payout", "Calculates net payout via dynamic pricing formula and triggers direct bank settlement before the farmer leaves.", PRIMARY),
        ("🤖 AI Krishi Mitra Advisory", "Multilingual voice/text assistant (Marathi/Hindi/English) + automated yield anomaly alerts for cattle health.", PRIMARY_DARK)
    ]

    for idx, (title, desc, col) in enumerate(pillars):
        row = idx // 2
        col_idx = idx % 2
        left_pos = 0.8 + col_idx * 5.95
        top_pos = 1.9 + row * 2.5
        card = add_card(slide3, left_pos, top_pos, 5.8, 2.3)
        tf = card.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = "\n" + desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MAIN

    # =========================================================================
    # SLIDE 4: HARDWARE TERMINAL & OPERATIONAL PIPELINE
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide4, DARK_BG)
    add_header(slide4, "HARDWARE-TO-CLOUD PIPELINE", "The 6-Step Automated Dairy Procurement Workflow", "From milk pouring to instant bank account credit in under 30 seconds", dark=True)

    steps = [
        ("Step 1", "Milk Weighing", "Weight streamed via Bluetooth BLE"),
        ("Step 2", "Quality Analysis", "Ultrasonic FAT & SNF analyzer"),
        ("Step 3", "Pricing Formula", "Dynamic rate calculation engine"),
        ("Step 4", "Dual Display Sync", "Customer screen live broadcast"),
        ("Step 5", "Instant Settlement", "UPI / IMPS 2-sec bank payout"),
        ("Step 6", "Passbook & Alert", "Mobile SMS / WhatsApp slip")
    ]

    for idx, (num, title, desc) in enumerate(steps):
        left_pos = 0.8 + idx * 1.98
        card = add_card(slide4, left_pos, 2.2, 1.85, 3.8, bg_color=RGBColor(30, 41, 59), border_color=PRIMARY)
        tf = card.text_frame
        tf.word_wrap = True
        
        p_n = tf.paragraphs[0]
        p_n.text = num.upper()
        p_n.font.size = Pt(10)
        p_n.font.bold = True
        p_n.font.color.rgb = ACCENT

        p_t = tf.add_paragraph()
        p_t.text = "\n" + title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = RGBColor(255, 255, 255)

        p_d = tf.add_paragraph()
        p_d.text = "\n" + desc
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = RGBColor(148, 163, 184)

    # Bottom summary banner
    banner = add_card(slide4, 0.8, 6.2, 11.7, 0.8, bg_color=RGBColor(20, 83, 45), border_color=RGBColor(34, 197, 94))
    tf_b = banner.text_frame
    p_ban = tf_b.paragraphs[0]
    p_ban.text = "⚡ Zero Operator Tampering • 100% Hardware Data Integrity • 2-Second Direct Farmer Credit"
    p_ban.font.size = Pt(13)
    p_ban.font.bold = True
    p_ban.font.color.rgb = RGBColor(220, 252, 231)
    p_ban.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 5: TECHNICAL ARCHITECTURE & HYBRID PERSISTENCE
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide5, LIGHT_BG)
    add_header(slide5, "SYSTEM ARCHITECTURE", "Enterprise-Grade, Scalable & Offline-First Tech Stack", "Built for high-concurrency rural connectivity resilience")

    tech_cards = [
        ("💻 Frontend & Presentation", "• Next.js 16 (App Router) & React 19\n• Tailwind CSS & Framer Motion\n• Web Bluetooth API (BLE 4.0+)\n• BroadcastChannel Multi-Screen Sync"),
        ("⚙️ Core Application Gateway", "• Dynamic Milk Pricing Formula Engine\n• Milk Drop Anomaly Detection Engine\n• Zero-Knowledge PII Masking (Aadhaar/Bank)\n• OTP-Guarded Profile Security Shield"),
        ("💾 Hybrid Database & Sync", "• Write-Through Sequential Lock Queue\n• Supabase Cloud PostgreSQL Sync\n• Local LowDB Edge Persistence Cache\n• Zero-loss offline rural fallback"),
        ("🏦 Fintech & AI Infrastructure", "• NPCI UPI / IMPS Instant Payout APIs\n• RazorpayX Sandbox & Live Gateway\n• Google Gemini AI Multilingual LLM\n• SMS / WhatsApp Transaction Dispatcher")
    ]

    for idx, (title, desc) in enumerate(tech_cards):
        row = idx // 2
        col_idx = idx % 2
        left_pos = 0.8 + col_idx * 5.95
        top_pos = 1.9 + row * 2.5
        card = add_card(slide5, left_pos, top_pos, 5.8, 2.3)
        tf = card.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY

        p2 = tf.add_paragraph()
        p2.text = "\n" + desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MAIN

    # =========================================================================
    # SLIDE 6: BUSINESS MODEL & MONETIZATION
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide6, LIGHT_BG)
    add_header(slide6, "BUSINESS & REVENUE MODEL", "How 3T Dairy Generates Sustainable, Scalable Cash Flow", "Multiple high-margin revenue streams across SaaS, Fintech & Agri-Commerce")

    streams = [
        ("🏢 1. B2B Dairy SaaS Subscription", 
         "Monthly software license charged per village milk collection hub.",
         "₹500 – ₹1,500 / Month per Center",
         "Recurring Annual Revenue"),
        ("⚡ 2. Instant Micro-Payout Fee", 
         "Nominal convenience fee per instant bank settlement transaction.",
         "₹0.50 – ₹1.00 per Payout",
         "High Volume Daily Micro-revenue"),
        ("🌾 3. Agri-Commerce & Cattle Feed", 
         "Marketplace commission on cattle feed, mineral mixtures, and vet supplies.",
         "3% – 5% Commission per Order",
         "Direct-to-farmer supplies"),
        ("🏦 4. Cattle Loan & Credit Scoring", 
         "Digitized milk ledger enables instant credit scoring for bank cattle loans.",
         "1% – 2% Loan Origination Fee",
         "Fintech credit enablement")
    ]

    for idx, (title, desc, rate, badge_text) in enumerate(streams):
        left_pos = 0.8 + idx * 2.95
        card = add_card(slide6, left_pos, 1.9, 2.85, 4.8)
        tf = card.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY

        p2 = tf.add_paragraph()
        p2.text = "\n" + desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MAIN

        p3 = tf.add_paragraph()
        p3.text = "\n\n💰 " + rate
        p3.font.size = Pt(12)
        p3.font.bold = True
        p3.font.color.rgb = PRIMARY_DARK

        p4 = tf.add_paragraph()
        p4.text = "\n📊 " + badge_text
        p4.font.size = Pt(10)
        p4.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 7: COMPETITIVE ADVANTAGE & NATIONAL IMPACT
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide7, LIGHT_BG)
    add_header(slide7, "COMPETITIVE MOAT & IMPACT", "Why 3T Outperforms Traditional & Existing Solutions", "Aligned with national development priorities and rural empowerment")

    # Table comparison
    table_shape = slide7.shapes.add_table(5, 4, Inches(0.8), Inches(1.9), Inches(11.7), Inches(2.8))
    table = table_shape.table
    table.columns[0].width = Inches(2.7)
    table.columns[1].width = Inches(3.0)
    table.columns[2].width = Inches(3.0)
    table.columns[3].width = Inches(3.0)

    headers = ["Feature", "Traditional Centers", "Legacy Software (Stellapps/Prompt)", "3T (Time To Time)"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY if i == 3 else DARK_BG
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)

    rows_data = [
        ("Payment Settlement", "15–20 Days Delay", "7–10 Days Batch", "✅ Instant 2-Second Direct UPI"),
        ("Testing Transparency", "Manual paper ledger", "Operator monitor only", "✅ Live Customer Dual-Display"),
        ("Hardware Integration", "Manual typing", "Proprietary locked hardware", "✅ Universal Web Bluetooth BLE"),
        ("AI & Health Monitoring", "None", "Basic historical charts", "✅ Multilingual AI Mitra + Drop Alert")
    ]

    for r_idx, row in enumerate(rows_data):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(240, 253, 244) if c_idx == 3 else CARD_BG
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                if c_idx == 3:
                    p.font.bold = True
                    p.font.color.rgb = PRIMARY_DARK
                else:
                    p.font.color.rgb = TEXT_MAIN

    # National alignment cards
    impacts = [
        ("👩‍🌾 Women Empowerment", "70%+ of Indian dairy labor is handled by rural women. Direct mobile UPI credits ensure earnings reach women directly."),
        ("🇮🇳 Digital India & UPI", "Deepens formal fintech penetration into Tier-3 & Tier-4 rural agrarian economies."),
        ("🥛 NDDB Clean Milk Mission", "Quality-based transparent pricing rewards hygienic milk production and penalizes adulteration.")
    ]
    for idx, (t, d) in enumerate(impacts):
        left_pos = 0.8 + idx * 3.95
        card = add_card(slide7, left_pos, 5.0, 3.8, 1.8)
        tf = card.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = t
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY
        p2 = tf.add_paragraph()
        p2.text = "\n" + d
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_MAIN

    # =========================================================================
    # SLIDE 8: ROADMAP & LIVE DEMO
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide8, DARK_BG)
    add_header(slide8, "EXECUTION ROADMAP & DEMO", "From Prototype to Pan-India Dairy Digitization", "Scalable 18-month execution plan", dark=True)

    phases = [
        ("Phase 1: Ground Pilot", "Q1 2026", "• 1 Village Hub (Sangamner)\n• 100+ Active Farmers\n• Real UPI payment testing\n• Hardware scale pairing"),
        ("Phase 2: District Scale", "Q2–Q3 2026", "• 50 Village Collection Centers\n• 5,000+ Daily Pours\n• Cattle Feed Marketplace\n• Seed grant application"),
        ("Phase 3: Cooperative Scale", "Q4 2026+", "• Integration with State Dairy Federations\n• 100,000+ Farmers\n• Formal NBFC Cattle Loans\n• National Expansion")
    ]

    for idx, (title, date_str, desc) in enumerate(phases):
        left_pos = 0.8 + idx * 3.95
        card = add_card(slide8, left_pos, 1.9, 3.8, 3.6, bg_color=RGBColor(30, 41, 59), border_color=PRIMARY)
        tf = card.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(134, 239, 172)

        p_d = tf.add_paragraph()
        p_d.text = date_str
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = ACCENT

        p2 = tf.add_paragraph()
        p2.text = "\n" + desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = RGBColor(226, 232, 240)

    # Demo Call-to-action
    cta = add_card(slide8, 0.8, 5.8, 11.7, 1.1, bg_color=RGBColor(20, 83, 45), border_color=RGBColor(34, 197, 94))
    tf_cta = cta.text_frame
    p_c1 = tf_cta.paragraphs[0]
    p_c1.text = "🚀 Experience the Live Working Prototype"
    p_c1.font.size = Pt(16)
    p_c1.font.bold = True
    p_c1.font.color.rgb = RGBColor(255, 255, 255)
    p_c1.alignment = PP_ALIGN.CENTER
    p_c2 = tf_cta.add_paragraph()
    p_c2.text = "Web: 3T Dairy Ecosystem | GitHub: sandeshk0019-spec/3tdairy | Contact: Sandesh Kadam"
    p_c2.font.size = Pt(11)
    p_c2.font.color.rgb = RGBColor(187, 247, 208)
    p_c2.alignment = PP_ALIGN.CENTER

    output_path = "3T_Dairy_Premium_Pitch_Deck.pptx"
    prs.save(output_path)
    print(f"Presentation generated successfully: {output_path}")

if __name__ == "__main__":
    create_presentation()
